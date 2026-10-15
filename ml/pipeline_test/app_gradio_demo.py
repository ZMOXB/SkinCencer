import gradio as gr
import tensorflow as tf
from tensorflow.keras.applications.resnet50 import preprocess_input
import numpy as np
from PIL import Image
import json
import os
import hashlib
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

IMG_SIZE = 224
MODEL_PATH = "Model3.keras"

# 1. โหลดโมเดล
print(f"กำลังโหลดโมเดลจาก {MODEL_PATH}...")
model = tf.keras.models.load_model(MODEL_PATH)
print(" โหลดโมเดลสำเร็จแล้ว!")

# 2. โหลดชื่อคลาสจาก class_indices.json
with open("class_indices.json", "r", encoding="utf-8") as f:
    class_indices = json.load(f)

index_to_class = {int(v): k for k, v in class_indices.items()}
class_names = [index_to_class[i] for i in sorted(index_to_class.keys())]

# คำแปลชื่อโรคภาษาไทย
disease_thai = {
    "Actinic Keratosis": "ระยะก่อนมะเร็งผิวหนัง (Actinic Keratosis)",
    "BCC": "มะเร็งผิวหนังเบเซลเซลล์ (BCC)",
    "Keratosis": "โรคกระเนื้อ (Keratosis)",
    "Melanocytic Nevi": "ไฝ หรือ ปานธรรมดา (Melanocytic Nevi)",
    "Melanoma": "มะเร็งผิวหนังเมลาโนมา (Melanoma)",
    "Psoriasis": "โรคสะเก็ดเงิน (Psoriasis)",
    "SCC": "มะเร็งผิวหนังชนิดสความัสเซลล์ (SCC)",
    "Tinea": "โรคเชื้อราที่ผิวหนัง / กลาก (Tinea)"
}

# ==========================================
# ฟังก์ชันช่วย: Center Crop พร้อมระบบ Zoom ซูมเข้า-ซูมออก
# ==========================================
def crop_with_zoom(img, zoom=1.0, pan_x=0, pan_y=0):
    if img is None:
        return None

    img_rgb = img.convert("RGB")
    w, h = img_rgb.size
    min_dim = min(w, h)
    crop_size = max(1, int(min_dim / max(zoom, 0.1)))

    # คำนวณจุดกึ่งกลาง (cx, cy) ตามการเลื่อนตำแหน่ง (pan_x, pan_y ระหว่าง -100 ถึง 100%)
    # pan_x: ลบ = ซ้าย, บวก = ขวา | pan_y: ลบ = ขึ้น, บวก = ลง
    cx = int(w / 2.0 + (pan_x / 100.0) * (w / 2.0))
    cy = int(h / 2.0 + (pan_y / 100.0) * (h / 2.0))

    x1 = cx - crop_size // 2
    y1 = cy - crop_size // 2
    x2 = x1 + crop_size
    y2 = y1 + crop_size

    # กรณีขอบกล่องครอปหลุดออกจากขอบรูป ให้สร้าง canvas เติมพิกเซลขอบภาพ
    pad_left = max(0, -x1)
    pad_top = max(0, -y1)
    pad_right = max(0, x2 - w)
    pad_bottom = max(0, y2 - h)

    if pad_left > 0 or pad_top > 0 or pad_right > 0 or pad_bottom > 0:
        new_w = w + pad_left + pad_right
        new_h = h + pad_top + pad_bottom
        edge_pixel = img_rgb.getpixel((0, 0))
        canvas = Image.new('RGB', (new_w, new_h), color=edge_pixel)
        canvas.paste(img_rgb, (pad_left, pad_top))
        cx_c = cx + pad_left
        cy_c = cy + pad_top
        box = (cx_c - crop_size // 2, cy_c - crop_size // 2, cx_c - crop_size // 2 + crop_size, cy_c - crop_size // 2 + crop_size)
        cropped = canvas.crop(box)
    else:
        cropped = img_rgb.crop((x1, y1, x2, y2))

    return cropped.resize((IMG_SIZE, IMG_SIZE), Image.Resampling.LANCZOS)


# ==========================================
# ฟังก์ชันช่วย: สร้างไฟล์ Excel (.xlsx) สไตล์เป๊ะตามตารางรายงาน
# ==========================================
def export_styled_excel(data_rows, class_name="test"):
    clean_name = class_name.replace(" ", "_") if class_name else "test"
    filename = f"test_results_{clean_name}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ผลการทดสอบ"

    headers = ["รูปที่", "คลาส", "ผลที่ทำนาย", "ความมั่นใจ", "ผลการประเมิน"]
    ws.append(headers)

    # Styles
    font_header = Font(name="Tahoma", size=11, bold=True)
    font_normal = Font(name="Tahoma", size=10)
    font_bold_white = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")

    fill_header = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_correct = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")  # สีเขียว
    fill_wrong = PatternFill(start_color="E53935", end_color="E53935", fill_type="solid")    # สีแดง

    align_center = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    # Header row
    for col_num in range(1, 6):
        cell = ws.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    # Data rows
    for r_idx, row_data in enumerate(data_rows, start=2):
        ws.append([
            row_data["รูปที่"],
            row_data["คลาส"],
            row_data["ผลที่ทำนาย"],
            row_data["ความมั่นใจ"],
            row_data["ผลการประเมิน"]
        ])

        for col_num in range(1, 6):
            c = ws.cell(row=r_idx, column=col_num)
            c.font = font_normal
            c.alignment = align_center
            c.border = thin_border

        # ไฮไลต์สีเขียว/แดงในช่องผลการประเมิน
        eval_cell = ws.cell(row=r_idx, column=5)
        if eval_cell.value == "ถูกต้อง":
            eval_cell.fill = fill_correct
            eval_cell.font = font_bold_white
        elif eval_cell.value == "ผิด":
            eval_cell.fill = fill_wrong
            eval_cell.font = font_bold_white

    # ความกว้างคอลัมน์ให้อ่านง่าย
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 22
    ws.column_dimensions['D'].width = 16
    ws.column_dimensions['E'].width = 16

    wb.save(filename)
    return filename


# ==========================================
# ตรรกะการทำงาน (Logic)
# ==========================================
def update_crop_preview(img, zoom, pan_x, pan_y):
    if img is None:
        return None
    return crop_with_zoom(img, zoom=zoom, pan_x=pan_x, pan_y=pan_y)

def zoom_in(current_zoom, img, pan_x, pan_y):
    new_zoom = round(float(current_zoom) + 0.05, 2)
    new_zoom = min(5.0, max(0.5, new_zoom))
    return new_zoom, crop_with_zoom(img, zoom=new_zoom, pan_x=pan_x, pan_y=pan_y)

def zoom_out(current_zoom, img, pan_x, pan_y):
    new_zoom = round(float(current_zoom) - 0.05, 2)
    new_zoom = min(5.0, max(0.5, new_zoom))
    return new_zoom, crop_with_zoom(img, zoom=new_zoom, pan_x=pan_x, pan_y=pan_y)

def zoom_reset(img, pan_x, pan_y):
    return 1.0, crop_with_zoom(img, zoom=1.0, pan_x=pan_x, pan_y=pan_y)

def move_pan_x(cur_x, delta, img, zoom, cur_y):
    new_x = max(-100, min(100, int(cur_x + delta)))
    return new_x, crop_with_zoom(img, zoom=zoom, pan_x=new_x, pan_y=cur_y)

def move_pan_y(cur_y, delta, img, zoom, cur_x):
    new_y = max(-100, min(100, int(cur_y + delta)))
    return new_y, crop_with_zoom(img, zoom=zoom, pan_x=cur_x, pan_y=new_y)

def reset_pan(img, zoom):
    return 0, 0, crop_with_zoom(img, zoom=zoom, pan_x=0, pan_y=0)

def get_image_hash(img):
    if img is None:
        return None
    try:
        small = img.copy()
        small.thumbnail((128, 128))
        return hashlib.md5(small.convert("RGB").tobytes()).hexdigest()
    except Exception:
        return None

def get_display_df(records):
    if not records:
        return pd.DataFrame(columns=["รูปที่", "คลาส", "ผลที่ทำนาย", "ความมั่นใจ", "ผลการประเมิน"])
    clean_records = [
        {
            "รูปที่": r["รูปที่"],
            "คลาส": r["คลาส"],
            "ผลที่ทำนาย": r["ผลที่ทำนาย"],
            "ความมั่นใจ": r["ความมั่นใจ"],
            "ผลการประเมิน": r["ผลการประเมิน"]
        }
        for r in records
    ]
    return pd.DataFrame(clean_records)

def predict_single(img, zoom, pan_x, pan_y, target_class, records):
    if img is None:
        return None, "⚠️ กรุณาอัปโหลดรูปภาพ", "-", "-", {}, records, get_display_df(records), build_summary_md(records), None

    img_cropped = crop_with_zoom(img, zoom=zoom, pan_x=pan_x, pan_y=pan_y)
    img_arr = np.array(img_cropped)
    img_arr = np.expand_dims(img_arr, axis=0)
    img_arr = preprocess_input(img_arr)

    predictions = model.predict(img_arr, verbose=0)[0]
    top_idx = np.argmax(predictions)
    pred_class = index_to_class[top_idx]
    conf = round(float(predictions[top_idx] * 100), 1)

    # ตรวจสอบว่าตรงกับคลาสที่เลือกหรือไม่
    is_correct = (pred_class.lower() == target_class.lower())
    eval_text = "✅ ถูกต้อง (ตรงตามคลาส)" if is_correct else f"❌ ผิด (ทายได้ {pred_class})"

    result_disease = f"{pred_class} - {disease_thai.get(pred_class, '')}"
    result_conf = f"{conf:.1f}%"

    pred_data = {
        "pred_class": pred_class,
        "conf": conf,
        "eval_status": "ถูกต้อง" if is_correct else "ผิด"
    }

    # ตรวจสอบว่าภาพนี้เคยบันทึกในตารางแล้วหรือไม่ (ถ้าเป็นภาพเดิม -> อัปเดตผลทำนายใหม่ลงตารางทันที)
    img_hash = get_image_hash(img)
    updated_records = list(records) if records else []
    found_idx = -1
    for i, r in enumerate(updated_records):
        if r.get("_img_hash") == img_hash:
            found_idx = i
            break

    excel_file = None
    if found_idx != -1:
        # เป็นภาพเดิม: นำผลทำนายใหม่ไปอัปเดตแทนค่าเดิมทันที
        row_num = updated_records[found_idx]["รูปที่"]
        updated_records[found_idx]["คลาส"] = target_class
        updated_records[found_idx]["ผลที่ทำนาย"] = pred_class
        updated_records[found_idx]["ความมั่นใจ"] = conf
        updated_records[found_idx]["ผลการประเมิน"] = "ถูกต้อง" if is_correct else "ผิด"

        excel_file = export_styled_excel(updated_records, class_name=target_class)
        summary_md = build_summary_md(updated_records, updated_msg=f"อัปเดตผลทำนายใหม่ของ รูปที่ {row_num} ในตารางเรียบร้อย ({pred_class} {conf}%)")
        gr.Info(f"อัปเดตผลทำนายใหม่ลงในตาราง (รูปที่ {row_num}) เรียบร้อยแล้วครับ")
    else:
        summary_md = build_summary_md(updated_records)
        if len(updated_records) > 0:
            excel_file = export_styled_excel(updated_records, class_name=target_class)

    df = get_display_df(updated_records)
    return img_cropped, result_disease, result_conf, eval_text, pred_data, updated_records, df, summary_md, excel_file


def save_record(records, pred_data, target_class, img, zoom, pan_x, pan_y):
    if img is None:
        gr.Warning("กรุณาอัปโหลดรูปภาพก่อนกดบันทึกครับ")
        return records, get_display_df(records), build_summary_md(records), None

    # คำนวณทำนายสดจากระยะซูมและตำแหน่งปัจจุบันเสมอ
    img_cropped = crop_with_zoom(img, zoom=zoom, pan_x=pan_x, pan_y=pan_y)
    img_arr = preprocess_input(np.expand_dims(np.array(img_cropped), axis=0))
    predictions = model.predict(img_arr, verbose=0)[0]
    top_idx = np.argmax(predictions)
    pred_class = index_to_class[top_idx]
    conf = round(float(predictions[top_idx] * 100), 1)
    is_correct = (pred_class.lower() == target_class.lower())
    eval_status = "ถูกต้อง" if is_correct else "ผิด"

    img_hash = get_image_hash(img)
    updated_records = list(records) if records else []

    # ตรวจสอบว่าเป็นภาพเดิมที่เคยบันทึกแล้วหรือไม่
    found_idx = -1
    for i, r in enumerate(updated_records):
        if r.get("_img_hash") == img_hash:
            found_idx = i
            break

    if found_idx != -1:
        # ภาพเดิม -> อัปเดตข้อมูลแถวเดิม
        row_num = updated_records[found_idx]["รูปที่"]
        updated_records[found_idx]["คลาส"] = target_class
        updated_records[found_idx]["ผลที่ทำนาย"] = pred_class
        updated_records[found_idx]["ความมั่นใจ"] = conf
        updated_records[found_idx]["ผลการประเมิน"] = eval_status
        gr.Info(f"บันทึกทับผลใหม่ลงใน รูปที่ {row_num} เรียบร้อยแล้วครับ")
    else:
        # ภาพใหม่ -> บันทึกเป็นรูปใหม่
        new_idx = len(updated_records) + 1
        new_row = {
            "รูปที่": new_idx,
            "คลาส": target_class,
            "ผลที่ทำนาย": pred_class,
            "ความมั่นใจ": conf,
            "ผลการประเมิน": eval_status,
            "_img_hash": img_hash
        }
        updated_records.append(new_row)
        gr.Success(f"บันทึก รูปที่ {new_idx} ลงตารางเรียบร้อยครับ")

    excel_file = export_styled_excel(updated_records, class_name=target_class)
    summary_md = build_summary_md(updated_records)
    df = get_display_df(updated_records)

    return updated_records, df, summary_md, excel_file


def undo_record(records, target_class):
    if not records or len(records) == 0:
        gr.Info("ยังไม่มีข้อมูลให้ลบครับ")
        return records, get_display_df(records), build_summary_md(records), None

    updated_records = list(records)[:-1]

    if len(updated_records) > 0:
        excel_file = export_styled_excel(updated_records, class_name=target_class)
    else:
        excel_file = None

    summary_md = build_summary_md(updated_records)
    df = get_display_df(updated_records)
    return updated_records, df, summary_md, excel_file


def clear_records():
    empty_records = []
    df = get_display_df([])
    summary_md = build_summary_md(empty_records)
    gr.Info("ล้างตารางข้อมูลเรียบร้อยแล้ว พร้อมเริ่มชุดใหม่")
    return empty_records, df, summary_md, None


def build_summary_md(records, updated_msg=""):
    total = len(records) if records else 0
    if total == 0:
        return (
            "### 📊 ความคืบหน้า: บันทึกแล้ว **0 / 30 รูป**\n"
            "*ยังไม่มีข้อมูลในตาราง กรุณาอัปโหลดรูปแล้วกดปุ่มบันทึกผล*"
        )

    correct = sum(1 for r in records if r["ผลการประเมิน"] == "ถูกต้อง")
    wrong = total - correct
    acc = (correct / total) * 100

    target_badge = ""
    if total == 30:
        target_badge = "\n\n🎉 **ยินดีด้วยครับ! บันทึกผลครบ 30 รูปตามเป้าหมายแล้ว ดาวน์โหลดไฟล์ Excel ไปใช้งานได้เลยครับ**"
    elif total > 30:
        target_badge = f"\n\n*(บันทึกเกิน 30 รูป: ปัจจุบัน {total} รูป)*"

    msg_line = f"\n\n> 💡 {updated_msg}" if updated_msg else ""

    return (
        f"### 📊 ความคืบหน้า: บันทึกแล้ว **{total} / 30 รูป**\n"
        f"- ✅ **ถูกต้อง:** {correct} รูป\n"
        f"- ❌ **ทายผิด:** {wrong} รูป\n"
        f"- 📈 **ความแม่นยำ (Accuracy): {acc:.2f}%**"
        f"{target_badge}{msg_line}"
    )


# ==========================================
# การจัดวาง UI (Gradio Blocks)
# ==========================================
custom_css = """
.status-card {
    background-color: #f8fafc;
    border: 1px solid #e2e8f0;
    border-radius: 8px;
    padding: 16px;
}
"""

with gr.Blocks(title="Skin Disease AI - Interactive Test Suite", theme=gr.themes.Soft(primary_hue="blue"), css=custom_css) as demo:
    # เก็บ State ของตารางและผลทำนายปัจจุบัน
    records_state = gr.State([])
    pred_data_state = gr.State({})

    gr.Markdown("# 🔬 ระบบทดสอบจำแนกโรคผิวหนังทีละรูป & บันทึกสร้างไฟล์ Excel")

    with gr.Row():
        # ==========================================
        # ฝั่งซ้ายมือ: อัปโหลดรูปภาพต้นฉบับ & เลือกโรค
        # ==========================================
        with gr.Column(scale=5):
            gr.Markdown("### 📸 ภาพต้นฉบับ")
            target_disease = gr.Dropdown(
                choices=class_names,
                value="Melanoma",
                label="🏷️ เฉลยโรคจริงของชุดนี้ (True Class)"
            )
            single_img = gr.Image(
                type="pil",
                label="อัปโหลดรูปภาพต้นฉบับที่นี่",
                height=450
            )

        # ==========================================
        # ฝั่งขวามือ: พรีวิว + ทำนาย + ซูม/เลื่อน + บันทึก
        # ==========================================
        with gr.Column(scale=7):
            gr.Markdown("### 🎯 ผลลัพธ์ & เครื่องมือปรับภาพ")

            # ส่วนพรีวิว 224x224 + กล่องผลทำนาย + ปุ่มทำนาย & บันทึก
            with gr.Row():
                with gr.Column(scale=5):
                    cropped_preview = gr.Image(
                        type="pil",
                        label="พรีวิว 224x224 (ตามที่ส่งเข้า AI จริง)",
                        height=200
                    )
                with gr.Column(scale=7):
                    pred_disease_box = gr.Textbox(
                        label="🎯 ผลการทำนายโรค (Top-1)",
                        interactive=False,
                        placeholder="กดปุ่ม 'ทดสอบทำนาย' เพื่อดูผล"
                    )
                    with gr.Row():
                        pred_conf_box = gr.Textbox(label=" ความมั่นใจ (%)", interactive=False, scale=1)
                        pred_eval_box = gr.Textbox(label=" ผลการประเมิน", interactive=False, scale=1)
                    with gr.Row():
                        btn_predict = gr.Button("🔍 ทดสอบทำนาย", variant="secondary", scale=1)
                        btn_save = gr.Button("💾 พอใจ ➡️ บันทึก", variant="primary", scale=1)

            # เครื่องมือซูมภาพ และ เครื่องมือเลื่อนตำแหน่งภาพ (จัดวางคู่กันกะทัดรัด)
            with gr.Row():
                # เครื่องมือซูม
                with gr.Column(scale=6):
                    zoom_slider = gr.Slider(
                        minimum=0.5,
                        maximum=5.0,
                        step=0.05,
                        value=1.0,
                        label="🔍 ซูม (0.5x - 5.0x)"
                    )
                    with gr.Row():
                        btn_zoom_minus = gr.Button("➖ -0.05", variant="secondary", min_width=0, size="sm", scale=1)
                        btn_zoom_reset = gr.Button("🔄 1.0x", variant="secondary", min_width=0, size="sm", scale=1)
                        btn_zoom_plus = gr.Button("➕ +0.05", variant="secondary", min_width=0, size="sm", scale=1)

                # เครื่องมือเลื่อนตำแหน่ง
                with gr.Column(scale=6):
                    with gr.Row():
                        pan_x = gr.Slider(minimum=-100, maximum=100, step=1, value=0, label="↔️ ซ้าย-ขวา", scale=1)
                        pan_y = gr.Slider(minimum=-100, maximum=100, step=1, value=0, label="↕️ ขึ้น-ลง", scale=1)
                    with gr.Row():
                        btn_pan_left = gr.Button("⬅️", min_width=0, size="sm")
                        btn_pan_right = gr.Button("➡️", min_width=0, size="sm")
                        btn_pan_up = gr.Button("⬆️", min_width=0, size="sm")
                        btn_pan_down = gr.Button("⬇️", min_width=0, size="sm")
                        btn_pan_center = gr.Button("🎯 กลาง", min_width=0, size="sm", variant="secondary")

            # กล่องสรุปผลสะสม & ดาวน์โหลด Excel
            with gr.Row():
                with gr.Column(scale=7):
                    summary_box = gr.Markdown(build_summary_md([]))
                with gr.Column(scale=5):
                    download_excel_file = gr.File(
                        label="📥 ดาวน์โหลดไฟล์ Excel (.xlsx)",
                        interactive=False,
                        height=90
                    )
                    with gr.Row():
                        btn_undo = gr.Button("↩️ ลบรูปล่าสุด", variant="stop", size="sm", min_width=0)
                        btn_clear = gr.Button("🗑️ ล้างตารางเริ่มใหม่", variant="stop", size="sm", min_width=0)

    gr.Markdown("---")
    gr.Markdown("### 📋 ตารางข้อมูลผลการทดสอบสะสม (ตรงรูปแบบสำหรับนำไปใส่เล่มรายงานวิจัย)")
    table_display = gr.Dataframe(
        headers=["รูปที่", "คลาส", "ผลที่ทำนาย", "ความมั่นใจ", "ผลการประเมิน"],
        label="ตารางบันทึกผล",
        interactive=False
    )

    # Event Handlers:
    # 1. พรีวิวภาพเมื่ออัปโหลดรูป เลื่อนซูม หรือเลื่อนตำแหน่ง
    single_img.change(
        fn=update_crop_preview,
        inputs=[single_img, zoom_slider, pan_x, pan_y],
        outputs=cropped_preview
    )
    zoom_slider.change(
        fn=update_crop_preview,
        inputs=[single_img, zoom_slider, pan_x, pan_y],
        outputs=cropped_preview
    )
    pan_x.change(
        fn=update_crop_preview,
        inputs=[single_img, zoom_slider, pan_x, pan_y],
        outputs=cropped_preview
    )
    pan_y.change(
        fn=update_crop_preview,
        inputs=[single_img, zoom_slider, pan_x, pan_y],
        outputs=cropped_preview
    )

    btn_zoom_minus.click(
        fn=zoom_out,
        inputs=[zoom_slider, single_img, pan_x, pan_y],
        outputs=[zoom_slider, cropped_preview]
    )
    btn_zoom_plus.click(
        fn=zoom_in,
        inputs=[zoom_slider, single_img, pan_x, pan_y],
        outputs=[zoom_slider, cropped_preview]
    )
    btn_zoom_reset.click(
        fn=zoom_reset,
        inputs=[single_img, pan_x, pan_y],
        outputs=[zoom_slider, cropped_preview]
    )

    # ปุ่มทิศทางเลื่อนตำแหน่งภาพ
    btn_pan_left.click(
        fn=lambda x, img, z, y: move_pan_x(x, -5, img, z, y),
        inputs=[pan_x, single_img, zoom_slider, pan_y],
        outputs=[pan_x, cropped_preview]
    )
    btn_pan_right.click(
        fn=lambda x, img, z, y: move_pan_x(x, 5, img, z, y),
        inputs=[pan_x, single_img, zoom_slider, pan_y],
        outputs=[pan_x, cropped_preview]
    )
    btn_pan_up.click(
        fn=lambda y, img, z, x: move_pan_y(y, -5, img, z, x),
        inputs=[pan_y, single_img, zoom_slider, pan_x],
        outputs=[pan_y, cropped_preview]
    )
    btn_pan_down.click(
        fn=lambda y, img, z, x: move_pan_y(y, 5, img, z, x),
        inputs=[pan_y, single_img, zoom_slider, pan_x],
        outputs=[pan_y, cropped_preview]
    )
    btn_pan_center.click(
        fn=reset_pan,
        inputs=[single_img, zoom_slider],
        outputs=[pan_x, pan_y, cropped_preview]
    )

    # 2. กดทำนายผล (ถ้าเป็นภาพเดิมที่เคยบันทึกแล้ว จะอัปเดตตารางและ Excel ด้วยผลใหม่ให้อัตโนมัติทันที)
    btn_predict.click(
        fn=predict_single,
        inputs=[single_img, zoom_slider, pan_x, pan_y, target_disease, records_state],
        outputs=[cropped_preview, pred_disease_box, pred_conf_box, pred_eval_box, pred_data_state, records_state, table_display, summary_box, download_excel_file]
    )

    # 3. กดบันทึกลงตาราง
    btn_save.click(
        fn=save_record,
        inputs=[records_state, pred_data_state, target_disease, single_img, zoom_slider, pan_x, pan_y],
        outputs=[records_state, table_display, summary_box, download_excel_file]
    )

    # 4. ลบรูปล่าสุด (Undo)
    btn_undo.click(
        fn=undo_record,
        inputs=[records_state, target_disease],
        outputs=[records_state, table_display, summary_box, download_excel_file]
    )

    # 5. ล้างตารางเริ่มใหม่
    btn_clear.click(
        fn=clear_records,
        inputs=[],
        outputs=[records_state, table_display, summary_box, download_excel_file]
    )

if __name__ == "__main__":
    demo.launch()